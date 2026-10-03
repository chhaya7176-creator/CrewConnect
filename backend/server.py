from fastapi import FastAPI, APIRouter, HTTPException, Depends, Header, UploadFile, File, Query, Response
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import io
import re
import json
import asyncio
import secrets
import base64
import logging
from pathlib import Path
from pydantic import BaseModel, Field, EmailStr
from typing import List, Optional
import uuid
import requests
import jwt
from datetime import datetime, timezone, timedelta

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

JWT_SECRET = os.environ.get('JWT_SECRET', 'dev_secret')
EMERGENT_LLM_KEY = os.environ.get('EMERGENT_LLM_KEY')
ADMIN_MOBILES = [m.strip() for m in os.environ.get('ADMIN_MOBILES', '').split(',') if m.strip()]

# ---------------- DPDP (India Digital Personal Data Protection Act) ----------------
ACTIVE_FILTER = {"deletion_requested": {"$ne": True}}  # exclude records with a pending DPDP deletion request
DPDP_CONSENT_VERSION = "1.0"
GRIEVANCE_OFFICER = {
    "name": "Gyanindra Mishra",
    "email": "gyanindramishra@gmail.com",
    "role": "Grievance Officer / Data Protection Contact",
}

# ---------------- Object storage ----------------
STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
APP_NAME = "maritime-connect"
storage_key = None
MIME_TYPES = {
    "jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "gif": "image/gif",
    "webp": "image/webp", "pdf": "application/pdf", "doc": "application/msword",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain",
}


def init_storage(force: bool = False):
    global storage_key
    if storage_key and not force:
        return storage_key
    resp = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_LLM_KEY}, timeout=30)
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key


def put_object(path: str, data: bytes, content_type: str) -> dict:
    key = init_storage()
    resp = requests.put(f"{STORAGE_URL}/objects/{path}",
                        headers={"X-Storage-Key": key, "Content-Type": content_type},
                        data=data, timeout=120)
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.put(f"{STORAGE_URL}/objects/{path}",
                            headers={"X-Storage-Key": key, "Content-Type": content_type},
                            data=data, timeout=120)
    resp.raise_for_status()
    return resp.json()


def get_object(path: str):
    key = init_storage()
    resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    resp.raise_for_status()
    return resp.content, resp.headers.get("Content-Type", "application/octet-stream")

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# ---------------- Config / Dropdown data ----------------
RANKS = [
    "Captain / Master", "Chief Officer", "Second Officer", "Third Officer",
    "Deck Cadet", "Chief Engineer", "Second Engineer", "Third Engineer",
    "Fourth Engineer", "Engine Cadet", "Electrical Officer (ETO)", "Bosun",
    "Able Seaman (AB)", "Ordinary Seaman (OS)", "Oiler", "Fitter",
    "Chief Cook", "Steward / Messman",
]
# Extra ranks available ONLY in the company Crew Planner (not in chat/requirement dropdowns)
CREW_PLANNER_RANKS = RANKS + [
    "Fourth Officer", "Junior Officer", "Assistant Electrical Officer (AETO)",
    "Trainee OS", "Trainee Wiper", "Trainee GS", "Trainee Fitter",
]
VESSEL_TYPES = [
    "Container", "Oil Tanker", "Chemical Tanker", "LNG / LPG Tanker",
    "Bulk Carrier", "General Cargo", "RoRo", "Passenger / Cruise",
    "Offshore / OSV", "Reefer", "Car Carrier",
]
SEA_TIME = ["0–9 months", "9–18 months", "18–30 months", "30+ months"]
COC_SEEKER = [
    "India", "Indonesia", "UK", "Singapore", "Australia", "Bangladesh",
    "Myanmar", "Sri Lanka", "Malaysia", "New Zealand", "Bahamas", "Panama",
    "Honduras", "Liberia", "Fiji", "Ukraine", "Others",
]
COC_PROVIDER = [
    "India", "Indonesia", "UK", "Singapore", "Australia", "Bangladesh",
    "Myanmar", "Sri Lanka", "Malaysia", "New Zealand", "Bahamas", "Panama",
    "Honduras", "Liberia", "Fiji", "Ethiopia", "Ukraine", "Any", "Others",
]
US_VISA_SEEKER = ["Valid", "Not Valid", "Not Available"]
US_VISA_PROVIDER = ["Must Have", "Preferred", "Not Required"]


@api_router.get("/config")
async def get_config():
    return {
        "ranks": RANKS,
        "crew_planner_ranks": CREW_PLANNER_RANKS,
        "vessel_types": VESSEL_TYPES,
        "sea_time": SEA_TIME,
        "coc_seeker": COC_SEEKER,
        "coc_provider": COC_PROVIDER,
        "us_visa_seeker": US_VISA_SEEKER,
        "us_visa_provider": US_VISA_PROVIDER,
        "grievance_officer": GRIEVANCE_OFFICER,
        "consent_version": DPDP_CONSENT_VERSION,
    }


# ---------------- Models ----------------
class OTPRequest(BaseModel):
    mobile: str


class OTPVerify(BaseModel):
    mobile: str
    otp: str
    role: Optional[str] = None
    ref: Optional[str] = None  # referral code the new user signed up through


SIGNUP_TARGET = 2  # real sign-ups through a referral link needed to auto-unlock the ₹499 boost


class SummaryRequest(BaseModel):
    role: str  # seeker | provider
    data: dict


class SeekerProfile(BaseModel):
    rank: str
    vessel_type: str
    sea_time: str
    coc: str
    us_visa: str
    other_visa: Optional[str] = ""
    mobile: str
    consent: bool
    name: str
    email: str
    availability: str
    summary: Optional[str] = ""
    boosted: bool = False


class ProviderRequirement(BaseModel):
    rank: str
    vessel_type: str
    sea_time: str
    sea_time_priority: Optional[str] = "Mandatory"
    coc: List[str]
    coc_priority: Optional[str] = "Mandatory"
    us_visa: str
    us_visa_priority: Optional[str] = "Preferred"
    mobile: str
    consent: bool
    company_name: str
    email: str
    availability: str
    availability_priority: Optional[str] = "Preferred"
    summary: Optional[str] = ""
    vessel_name: Optional[str] = ""
    port: Optional[str] = ""
    rpsl_number: Optional[str] = ""


class CrewNeedIn(BaseModel):
    vessel_name: str
    ranks: List[str]
    by_date: str
    port: str
    note: Optional[str] = ""


class FeedbackIn(BaseModel):
    category: str
    message: str


class AlertOptIn(BaseModel):
    enabled: bool
    whatsapp: Optional[str] = ""


class SavedSearchIn(BaseModel):
    name: str
    criteria: dict


# ---------------- Auth helpers ----------------
def mobile_key(m: str) -> str:
    """Canonical identity for a phone number: last 10 digits (ignores +, spaces, country code)."""
    digits = "".join(c for c in (m or "") if c.isdigit())
    return digits[-10:] if len(digits) >= 10 else digits


def to_e164(m: str) -> str:
    """Best-effort E.164 for display in the phone field (India-first, keeps existing '+')."""
    m = (m or "").strip()
    if m.startswith("+"):
        return m
    digits = "".join(c for c in m if c.isdigit())
    if not digits:
        return ""
    if len(digits) == 10:
        return "+91" + digits
    return "+" + digits


def create_token(user_id: str, mobile: str) -> str:
    payload = {
        "user_id": user_id,
        "mobile": mobile,
        "exp": datetime.now(timezone.utc) + timedelta(days=30),
    }
    return jwt.encode(payload, JWT_SECRET, algorithm="HS256")


async def get_current_user(authorization: Optional[str] = Header(None)):
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Not authenticated")
    token = authorization.split(" ", 1)[1]
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=["HS256"])
    except jwt.PyJWTError:
        raise HTTPException(status_code=401, detail="Invalid token")
    user = await db.users.find_one({"id": payload["user_id"]}, {"_id": 0})
    if not user:
        raise HTTPException(status_code=401, detail="User not found")
    user["is_admin"] = user.get("mobile") in ADMIN_MOBILES
    return user


async def require_admin(user=Depends(get_current_user)):
    if not user.get("is_admin"):
        raise HTTPException(status_code=403, detail="Admin access required")
    return user


# ---------------- Auth endpoints ----------------
@api_router.post("/auth/request-otp")
async def request_otp(req: OTPRequest):
    digits = "".join(c for c in req.mobile if c.isdigit())
    if len(digits) < 10 or len(digits) > 15:
        raise HTTPException(status_code=400, detail="Mobile number must be 10–15 digits")
    otp = "123456"  # simulated OTP for demo
    key = mobile_key(req.mobile)
    await db.otps.update_one(
        {"mobile_key": key},
        {"$set": {"mobile_key": key, "mobile": req.mobile, "otp": otp,
                  "created_at": datetime.now(timezone.utc).isoformat()}},
        upsert=True,
    )
    # In demo mode we return the OTP so the UI can show it
    return {"success": True, "demo_otp": otp, "message": "OTP sent (demo mode)"}


async def attribute_referral(ref_code: str, new_user_id: str):
    """Attribute a REAL new sign-up to the seafarer whose referral link was used.
    Counts one sign-up per new user and auto-unlocks the referrer's ₹499 boost at the target."""
    if not ref_code:
        return
    seeker = await db.seekers.find_one({"referral_code": ref_code}, {"_id": 0, "user_id": 1})
    if not seeker:
        return
    referrer_id = seeker.get("user_id")
    if not referrer_id or referrer_id == new_user_id:
        return  # no self-referrals
    await db.users.update_one({"id": referrer_id}, {"$inc": {"ref_signups": 1}})
    ref_user = await db.users.find_one({"id": referrer_id}, {"_id": 0}) or {}
    if ref_user.get("ref_signups", 0) >= SIGNUP_TARGET and not ref_user.get("ref_completed"):
        await db.users.update_one({"id": referrer_id}, {"$set": {"ref_completed": True}})
        await db.seekers.update_many({"user_id": referrer_id}, {"$set": {"boosted": True}})
        for s in await db.seekers.find({"user_id": referrer_id}, {"_id": 0}).to_list(100):
            await notify_companies_of_seeker(s)


@api_router.post("/auth/verify-otp")
async def verify_otp(req: OTPVerify):
    key = mobile_key(req.mobile)
    record = await db.otps.find_one({"mobile_key": key}, {"_id": 0}) \
        or await db.otps.find_one({"mobile": req.mobile}, {"_id": 0})
    if not record or record["otp"] != req.otp:
        raise HTTPException(status_code=400, detail="Invalid OTP")
    # Look up user by canonical key first, then fall back to legacy exact-mobile match.
    user = await db.users.find_one({"mobile_key": key}, {"_id": 0}) \
        or await db.users.find_one({"mobile": req.mobile}, {"_id": 0})
    if not user:
        user = {
            "id": str(uuid.uuid4()),
            "mobile": req.mobile,
            "mobile_key": key,
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
        if req.ref:
            user["referred_by"] = req.ref
        await db.users.insert_one(dict(user))
        if req.ref:
            await attribute_referral(req.ref, user["id"])
    elif not user.get("mobile_key"):
        await db.users.update_one({"id": user["id"]}, {"$set": {"mobile_key": key}})
        user["mobile_key"] = key
    token = create_token(user["id"], user["mobile"])
    if req.role in ("seeker", "provider"):
        await db.users.update_one({"id": user["id"]}, {"$set": {"role": req.role}})
        user["role"] = req.role
    user["is_admin"] = user["mobile"] in ADMIN_MOBILES
    return {"token": token, "user": user}


@api_router.get("/auth/me")
async def me(user=Depends(get_current_user)):
    return user


# ---------------- AI summary ----------------
async def generate_ai_summary(role: str, data: dict) -> str:
    fallback = build_template_summary(role, data)
    if not EMERGENT_LLM_KEY:
        return fallback
    try:
        from emergentintegrations.llm.chat import LlmChat, UserMessage
        if role == "seeker":
            sys = ("You are a maritime crewing assistant. Write a crisp, professional 2-3 sentence "
                   "profile summary for a seafarer job seeker based on the JSON data. Highlight rank, "
                   "vessel experience, sea time, certification and availability. Do not invent facts. "
                   "Do not include the mobile number. Keep under 60 words.")
        else:
            sys = ("You are a maritime crewing assistant. Write a crisp, professional 2-3 sentence "
                   "crew requirement summary for a shipping company based on the JSON data. Highlight "
                   "the required rank, vessel type, experience, certification preferences and joining "
                   "timeline. Do not invent facts. Keep under 60 words.")
        chat = LlmChat(api_key=EMERGENT_LLM_KEY, session_id=str(uuid.uuid4()), system_message=sys).with_model("openai", "gpt-5.4")
        msg = UserMessage(text=f"Data: {data}")
        resp = await asyncio.wait_for(chat.send_message(msg), timeout=12)
        text = resp if isinstance(resp, str) else str(resp)
        return text.strip() or fallback
    except Exception as e:
        logger.warning(f"AI summary failed: {e}")
        return fallback


def build_template_summary(role: str, d: dict) -> str:
    if role == "seeker":
        return (f"{d.get('rank','Seafarer')} with {d.get('sea_time','')} of sea time on "
                f"{d.get('vessel_type','')} vessels. Holds {d.get('coc','')} COC, US Visa status: "
                f"{d.get('us_visa','')}. Available from {d.get('availability','')}.")
    coc = d.get('coc')
    coc = ", ".join(coc) if isinstance(coc, list) else coc
    return (f"Seeking a {d.get('rank','crew')} for {d.get('vessel_type','')} vessels with "
            f"{d.get('sea_time','')} experience. Preferred COC: {coc}. US Visa: {d.get('us_visa','')}. "
            f"Joining: {d.get('availability','')}.")


@api_router.post("/ai/summary")
async def ai_summary(req: SummaryRequest):
    summary = await generate_ai_summary(req.role, req.data)
    return {"summary": summary}


# ---------------- Profile submission ----------------
@api_router.post("/seeker/profile")
async def create_seeker(profile: SeekerProfile, user=Depends(get_current_user)):
    doc = profile.model_dump()
    doc["id"] = str(uuid.uuid4())
    doc["user_id"] = user["id"]
    doc["role"] = "seeker"
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    doc["referral_code"] = uuid.uuid4().hex[:8]
    await db.seekers.insert_one(dict(doc))
    doc.pop("_id", None)
    await db.users.update_one({"id": user["id"]}, {"$set": {
        "consent_given": bool(profile.consent), "consent_at": datetime.now(timezone.utc).isoformat(),
        "consent_version": DPDP_CONSENT_VERSION}})
    return doc


@api_router.post("/provider/requirement")
async def create_provider(req: ProviderRequirement, user=Depends(get_current_user)):
    doc = req.model_dump()
    doc["id"] = str(uuid.uuid4())
    doc["user_id"] = user["id"]
    doc["role"] = "provider"
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    doc["referral_code"] = uuid.uuid4().hex[:8]
    await db.providers.insert_one(dict(doc))
    doc.pop("_id", None)
    await db.users.update_one({"id": user["id"]}, {"$set": {
        "consent_given": bool(req.consent), "consent_at": datetime.now(timezone.utc).isoformat(),
        "consent_version": DPDP_CONSENT_VERSION}})
    await notify_company_of_boosted_matches(doc)
    return doc


@api_router.put("/seeker/profile/{profile_id}")
async def update_seeker(profile_id: str, profile: SeekerProfile, user=Depends(get_current_user)):
    existing = await db.seekers.find_one({"id": profile_id, "user_id": user["id"]}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Profile not found")
    updates = profile.model_dump()
    updates.pop("boosted", None)  # preserve promotion state across edits
    await db.seekers.update_one({"id": profile_id}, {"$set": updates})
    doc = await db.seekers.find_one({"id": profile_id}, {"_id": 0})
    if doc.get("boosted"):
        await notify_companies_of_seeker(doc)
    return doc


@api_router.put("/provider/requirement/{req_id}")
async def update_provider(req_id: str, req: ProviderRequirement, user=Depends(get_current_user)):
    existing = await db.providers.find_one({"id": req_id, "user_id": user["id"]}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Requirement not found")
    await db.providers.update_one({"id": req_id}, {"$set": req.model_dump()})
    doc = await db.providers.find_one({"id": req_id}, {"_id": 0})
    await notify_company_of_boosted_matches(doc)
    return doc


# ---------------- Matching ----------------
def sea_time_rank(s: str) -> int:
    return {"0–9 months": 0, "9–18 months": 1, "18–30 months": 2, "30+ months": 3}.get(s, 0)


def matches_requirement(seeker: dict, req: dict) -> bool:
    if seeker["rank"] != req["rank"]:
        return False
    if seeker["vessel_type"] != req["vessel_type"] and req["vessel_type"] not in ("Any",):
        return False
    if req.get("sea_time_priority", "Mandatory") == "Mandatory":
        if sea_time_rank(seeker["sea_time"]) < sea_time_rank(req["sea_time"]):
            return False
    coc_list = req.get("coc", [])
    if req.get("coc_priority", "Mandatory") == "Mandatory":
        if "Any" not in coc_list and seeker["coc"] not in coc_list:
            return False
    return True


def criteria_status(seeker: dict, req: dict) -> dict:
    coc_list = req.get("coc", [])
    return {
        "Rank": seeker["rank"] == req["rank"],
        "Vessel": (seeker["vessel_type"] == req["vessel_type"]) or (req["vessel_type"] == "Any"),
        "Sea time": sea_time_rank(seeker["sea_time"]) >= sea_time_rank(req["sea_time"]),
        "COC": ("Any" in coc_list) or (seeker.get("coc") in coc_list),
    }


def matched_labels(seeker: dict, req: dict) -> list:
    return [k for k, v in criteria_status(seeker, req).items() if v]


def gap_labels(seeker: dict, req: dict) -> list:
    return [k for k, v in criteria_status(seeker, req).items() if not v]


def dedup_by(items: list, keys: list) -> list:
    seen, out = set(), []
    for it in items:
        sig = tuple(str(it.get(k)) for k in keys)
        if sig in seen:
            continue
        seen.add(sig)
        out.append(it)
    return out


def redact_seeker(seeker: dict, unlocked: bool) -> dict:
    out = {
        "id": seeker["id"],
        "rank": seeker["rank"],
        "vessel_type": seeker["vessel_type"],
        "sea_time": seeker["sea_time"],
        "coc": seeker["coc"],
        "us_visa": seeker["us_visa"],
        "other_visa": seeker.get("other_visa", ""),
        "availability": seeker["availability"],
        "summary": seeker.get("summary", ""),
        "boosted": seeker.get("boosted", False),
        "unlocked": unlocked,
    }
    if unlocked:
        out.update({"name": seeker["name"], "mobile": seeker["mobile"], "email": seeker["email"]})
    else:
        out.update({"name": "🔒 Locked", "mobile": "🔒 Locked", "email": "🔒 Locked"})
    return out


def is_eligible(user: dict) -> bool:
    return bool(user.get("reg_paid") or user.get("ref_completed"))


# ---------------- WhatsApp (Twilio) company alerts ----------------
TWILIO_ACCOUNT_SID = os.environ.get("TWILIO_ACCOUNT_SID")
TWILIO_AUTH_TOKEN = os.environ.get("TWILIO_AUTH_TOKEN")
TWILIO_WHATSAPP_FROM = os.environ.get("TWILIO_WHATSAPP_FROM")
TWILIO_MESSAGING_SERVICE_SID = os.environ.get("TWILIO_MESSAGING_SERVICE_SID")
TWILIO_CONTENT_SID = os.environ.get("TWILIO_CONTENT_SID")


def twilio_enabled() -> bool:
    return bool(TWILIO_ACCOUNT_SID and TWILIO_AUTH_TOKEN
                and (TWILIO_WHATSAPP_FROM or TWILIO_MESSAGING_SERVICE_SID))


async def is_subscribed(user_id: str) -> bool:
    """True if the company holds an active (non-expired) annual subscription."""
    u = await db.users.find_one({"id": user_id}, {"_id": 0}) or {}
    if not u.get("subscribed"):
        return False
    exp = u.get("subscription_expiry")
    if exp and exp < datetime.now(timezone.utc).isoformat():
        return False
    return True


async def company_is_paying(user_id: str) -> bool:
    """A company qualifies for proactive alerts once it is an engaged/paying customer:
    it has unlocked at least one CV, made a paid unlock, or holds a subscription."""
    u = await db.users.find_one({"id": user_id}, {"_id": 0}) or {}
    if await is_subscribed(user_id):
        return True
    if await db.unlocks.count_documents({"user_id": user_id}) > 0:
        return True
    if await db.payments.count_documents({"user_id": user_id, "type": "cv_unlock"}) > 0:
        return True
    return False


async def _twilio_send(to_e164: str, rank: str) -> dict:
    body = f"A new matching {rank} is available on Smart Crewing."

    def send_sync():
        from twilio.rest import Client
        c = Client(TWILIO_ACCOUNT_SID, TWILIO_AUTH_TOKEN)
        kwargs = {"to": f"whatsapp:{to_e164}"}
        if TWILIO_CONTENT_SID:
            kwargs["content_sid"] = TWILIO_CONTENT_SID
            kwargs["content_variables"] = json.dumps({"1": rank})
        else:
            kwargs["body"] = body
        if TWILIO_MESSAGING_SERVICE_SID:
            kwargs["messaging_service_sid"] = TWILIO_MESSAGING_SERVICE_SID
        else:
            kwargs["from_"] = TWILIO_WHATSAPP_FROM
        return c.messages.create(**kwargs)

    msg = await asyncio.to_thread(send_sync)
    return {"status": getattr(msg, "status", "queued"), "sid": getattr(msg, "sid", None)}


async def send_company_whatsapp_alert(*, company_user_id: str, to_mobile: str, rank: str,
                                      company_name: str, seeker_id: str):
    """Notify a paying company that a matching crew profile is now available.
    Deduped per (company_user_id, seeker_id). Runs in DEMO mode (no send, no charge) until
    Twilio credentials are configured, then sends via Twilio automatically."""
    to_wa = to_e164(to_mobile)
    if not to_wa:
        return
    if await db.wa_alerts.find_one({"company_user_id": company_user_id, "seeker_id": seeker_id}):
        return
    base = {"company_user_id": company_user_id, "seeker_id": seeker_id, "to": to_wa,
            "rank": rank, "company_name": company_name, "kind": "matching_crew",
            "message": f"A new matching {rank} is available on Smart Crewing.",
            "created_at": datetime.now(timezone.utc).isoformat()}
    if not twilio_enabled():
        await db.wa_alerts.insert_one({**base, "status": "simulated", "mode": "demo"})
        logger.info(f"WhatsApp DEMO alert -> {company_name} (matching {rank})")
        return
    try:
        res = await _twilio_send(to_wa, rank)
        await db.wa_alerts.insert_one({**base, "status": res["status"], "mode": "live", "sid": res["sid"]})
    except Exception as e:
        await db.wa_alerts.insert_one({**base, "status": "error", "mode": "live"})
        logger.warning(f"Twilio WhatsApp send failed: {e}")


async def notify_companies_of_seeker(seeker: dict):
    """Trigger A: a seeker's profile just went live/boosted -> alert matching paying companies."""
    if not seeker:
        return
    providers = await db.providers.find(ACTIVE_FILTER, {"_id": 0}).to_list(2000)
    for p in providers:
        if not matches_requirement(seeker, p):
            continue
        if not await company_is_paying(p["user_id"]):
            continue
        await send_company_whatsapp_alert(
            company_user_id=p["user_id"], to_mobile=p.get("mobile", ""),
            rank=seeker["rank"], company_name=p.get("company_name", "Company"),
            seeker_id=seeker["id"])


async def notify_company_of_boosted_matches(provider: dict):
    """Trigger B: a company became paying or posted/edited a matching vacancy ->
    alert that company about currently-available boosted crew."""
    if not provider or not await company_is_paying(provider["user_id"]):
        return
    seekers = await db.seekers.find({"boosted": True, **ACTIVE_FILTER}, {"_id": 0}).to_list(2000)
    for s in seekers:
        if not matches_requirement(s, provider):
            continue
        await send_company_whatsapp_alert(
            company_user_id=provider["user_id"], to_mobile=provider.get("mobile", ""),
            rank=s["rank"], company_name=provider.get("company_name", "Company"),
            seeker_id=s["id"])


def provider_card(p: dict, elig: bool) -> dict:
    return {
        "id": p["id"],
        "company_name": p["company_name"] if elig else "🔒 Locked",
        "company_mobile": p.get("mobile", "") if elig else "",
        "company_email": p.get("email", "") if elig else "",
        "rank": p["rank"],
        "vessel_type": p["vessel_type"],
        "sea_time": p["sea_time"],
        "sea_time_priority": p.get("sea_time_priority", "Mandatory"),
        "coc": p["coc"],
        "coc_priority": p.get("coc_priority", "Mandatory"),
        "us_visa": p["us_visa"],
        "us_visa_priority": p.get("us_visa_priority", "Preferred"),
        "availability": p["availability"],
        "availability_priority": p.get("availability_priority", "Preferred"),
        "summary": p.get("summary", ""),
        "vessel_name": p.get("vessel_name", ""),
        "port": p.get("port", ""),
    }


@api_router.get("/matches")
async def get_matches(user=Depends(get_current_user)):
    seeker = await db.seekers.find_one({"user_id": user["id"]}, {"_id": 0}, sort=[("created_at", -1)])
    provider = await db.providers.find_one({"user_id": user["id"]}, {"_id": 0}, sort=[("created_at", -1)])

    result = {"role": None, "matches": []}

    if provider:
        result["role"] = "provider"
        unlocks = await db.unlocks.find({"user_id": user["id"]}, {"_id": 0}).to_list(1000)
        unlocked_ids = {u["seeker_id"] for u in unlocks}
        seekers = await db.seekers.find(ACTIVE_FILTER, {"_id": 0}).to_list(1000)
        matched, near = [], []
        for s in seekers:
            if matches_requirement(s, provider):
                matched.append(s)
            elif s["rank"] == provider["rank"]:
                near.append(s)
        matched.sort(key=lambda s: (not s.get("boosted", False), s.get("created_at", "")), reverse=False)
        near = dedup_by(near, ["rank", "vessel_type", "sea_time", "coc", "us_visa"])
        seen_at = user.get("matches_seen_at") or user.get("created_at") or ""
        result["new_matches"] = sum(1 for s in matched if (s.get("created_at") or "") > seen_at)
        result["matches"] = [{**redact_seeker(s, s["id"] in unlocked_ids),
                              "match_reasons": matched_labels(s, provider)} for s in matched]
        result["near_matches"] = [{**redact_seeker(s, False),
                                   "gaps": gap_labels(s, provider)} for s in near[:6]]
        result["free_unlocks_left"] = max(0, FREE_UNLOCKS - len(unlocks))
        result["requirement"] = provider
    elif seeker:
        result["role"] = "seeker"
        elig = is_eligible(user)
        result["eligible"] = elig
        providers = await db.providers.find(ACTIVE_FILTER, {"_id": 0}).to_list(1000)
        matched, near = [], []
        for p in providers:
            if matches_requirement(seeker, p):
                matched.append(p)
            elif p["rank"] == seeker["rank"]:
                near.append(p)
        seen_at = user.get("matches_seen_at") or user.get("created_at") or ""
        result["new_matches"] = sum(1 for p in matched if (p.get("created_at") or "") > seen_at)
        result["matches"] = [{**provider_card(p, elig),
                              "match_reasons": matched_labels(seeker, p)} for p in matched]
        near = dedup_by(near, ["company_name", "rank", "vessel_type", "sea_time"])
        result["near_matches"] = [{**provider_card(p, elig),
                                   "gaps": gap_labels(seeker, p)} for p in near[:6]]
        result["profile"] = seeker
    return result


@api_router.post("/matches/mark-seen")
async def mark_matches_seen(user=Depends(get_current_user)):
    await db.users.update_one({"id": user["id"]}, {"$set": {"matches_seen_at": datetime.now(timezone.utc).isoformat()}})
    return {"ok": True}


@api_router.get("/alerts/status")
async def alerts_status(user=Depends(get_current_user)):
    return {"enabled": bool(user.get("alert_optin")), "whatsapp": to_e164(user.get("alert_whatsapp") or user.get("mobile", ""))}


@api_router.post("/alerts/optin")
async def alerts_optin(body: AlertOptIn, user=Depends(get_current_user)):
    wa = to_e164(body.whatsapp or user.get("mobile", ""))
    await db.users.update_one({"id": user["id"]}, {"$set": {"alert_optin": body.enabled, "alert_whatsapp": wa}})
    return {"ok": True, "enabled": body.enabled, "whatsapp": wa}


# ---------------- Saved searches (companies) ----------------
def _search_req(criteria: dict) -> dict:
    return {
        "rank": criteria.get("rank"),
        "vessel_type": criteria.get("vessel_type", "Any"),
        "sea_time": criteria.get("sea_time", "0–9 months"),
        "sea_time_priority": criteria.get("sea_time_priority", "Mandatory"),
        "coc": criteria.get("coc", ["Any"]) or ["Any"],
        "coc_priority": criteria.get("coc_priority", "Mandatory"),
    }


@api_router.post("/saved-searches")
async def create_saved_search(body: SavedSearchIn, user=Depends(get_current_user)):
    provider = await db.providers.find_one({"user_id": user["id"]}, {"_id": 0})
    if not provider:
        raise HTTPException(status_code=403, detail="Only companies can save searches")
    now = datetime.now(timezone.utc).isoformat()
    doc = {
        "id": str(uuid.uuid4()),
        "user_id": user["id"],
        "name": body.name or body.criteria.get("rank", "Search"),
        "criteria": body.criteria,
        "created_at": now,
        "last_viewed_at": now,
    }
    await db.saved_searches.insert_one(dict(doc))
    doc.pop("_id", None)
    return doc


@api_router.get("/saved-searches")
async def list_saved_searches(user=Depends(get_current_user)):
    searches = await db.saved_searches.find({"user_id": user["id"]}, {"_id": 0}).sort("created_at", -1).to_list(100)
    seekers = await db.seekers.find(ACTIVE_FILTER, {"_id": 0}).to_list(1000)
    out = []
    for s in searches:
        req = _search_req(s["criteria"])
        matched = [k for k in seekers if matches_requirement(k, req)]
        seen_at = s.get("last_viewed_at") or ""
        out.append({**s, "count": len(matched),
                    "new_count": sum(1 for k in matched if (k.get("created_at") or "") > seen_at)})
    return {"searches": out}


@api_router.get("/saved-searches/{sid}/results")
async def saved_search_results(sid: str, user=Depends(get_current_user)):
    s = await db.saved_searches.find_one({"id": sid, "user_id": user["id"]}, {"_id": 0})
    if not s:
        raise HTTPException(status_code=404, detail="Search not found")
    unlocks = await db.unlocks.find({"user_id": user["id"]}, {"_id": 0}).to_list(1000)
    unlocked_ids = {u["seeker_id"] for u in unlocks}
    req = _search_req(s["criteria"])
    seekers = await db.seekers.find(ACTIVE_FILTER, {"_id": 0}).to_list(1000)
    matched = [k for k in seekers if matches_requirement(k, req)]
    matched.sort(key=lambda k: k.get("created_at", ""), reverse=True)  # newest first
    seen_at = s.get("last_viewed_at") or ""
    results = [{**redact_seeker(k, k["id"] in unlocked_ids),
                "match_reasons": matched_labels(k, req),
                "is_new": (k.get("created_at") or "") > seen_at} for k in matched]
    await db.saved_searches.update_one({"id": sid}, {"$set": {"last_viewed_at": datetime.now(timezone.utc).isoformat()}})
    return {"search": s, "results": results, "free_unlocks_left": max(0, FREE_UNLOCKS - len(unlocks))}


@api_router.delete("/saved-searches/{sid}")
async def delete_saved_search(sid: str, user=Depends(get_current_user)):
    await db.saved_searches.delete_one({"id": sid, "user_id": user["id"]})
    return {"ok": True}


@api_router.get("/browse")
async def browse_all(user=Depends(get_current_user)):
    provider = await db.providers.find_one({"user_id": user["id"]}, {"_id": 0}, sort=[("created_at", -1)])
    seeker = await db.seekers.find_one({"user_id": user["id"]}, {"_id": 0}, sort=[("created_at", -1)])
    result = {"role": None, "matches": []}
    urole = user.get("role")
    is_company = (urole == "provider") if urole in ("provider", "seeker") else (bool(provider) and not seeker)
    if is_company:
        result["role"] = "provider"
        unlocks = await db.unlocks.find({"user_id": user["id"]}, {"_id": 0}).to_list(2000)
        unlocked_ids = {u["seeker_id"] for u in unlocks}
        seekers = await db.seekers.find(ACTIVE_FILTER, {"_id": 0}).to_list(2000)
        seekers.sort(key=lambda s: (not s.get("boosted", False), s.get("created_at", "")))
        result["matches"] = [redact_seeker(s, s["id"] in unlocked_ids) for s in seekers]
        result["free_unlocks_left"] = max(0, FREE_UNLOCKS - len(unlocks))
        result["requirement"] = provider
    else:
        result["role"] = "seeker"
        elig = is_eligible(user)
        result["eligible"] = elig
        providers = await db.providers.find(ACTIVE_FILTER, {"_id": 0}).to_list(2000)
        result["matches"] = [provider_card(p, elig) for p in providers]
    return result


class UnlockRequest(BaseModel):
    seeker_id: str


FREE_UNLOCKS = 5  # first 5 matched crew contacts free per company, then ₹29 each
UNLOCK_PRICE = 29
SUBSCRIPTION_PRICE = 199000


async def _do_unlock(user_id: str, seeker_id: str, amount: int):
    now = datetime.now(timezone.utc).isoformat()
    await db.unlocks.update_one(
        {"user_id": user_id, "seeker_id": seeker_id},
        {"$set": {"user_id": user_id, "seeker_id": seeker_id, "amount": amount, "at": now}},
        upsert=True,
    )
    if amount:
        await db.payments.insert_one({"user_id": user_id, "type": "cv_unlock", "amount": amount, "at": now})
    provider = await db.providers.find_one({"user_id": user_id}, {"_id": 0}, sort=[("created_at", -1)])
    if provider:
        await notify_company_of_boosted_matches(provider)


@api_router.post("/unlock-cv")
async def unlock_cv(req: UnlockRequest, user=Depends(get_current_user)):
    """First 5 crew contacts are free (unlimited while subscribed); after that a ₹29 payment
    is required — the client runs the Razorpay flow, then completes the unlock via /pay/verify."""
    seeker = await db.seekers.find_one({"id": req.seeker_id}, {"_id": 0})
    if not seeker:
        raise HTTPException(status_code=404, detail="Profile not found")
    used = await db.unlocks.count_documents({"user_id": user["id"]})
    existing = await db.unlocks.find_one({"user_id": user["id"], "seeker_id": req.seeker_id}, {"_id": 0})
    if existing:
        return {"success": True, "amount": existing.get("amount", 0), "already": True,
                "free_unlocks_left": max(0, FREE_UNLOCKS - used), "cv": redact_seeker(seeker, True)}
    subscribed = await is_subscribed(user["id"])
    if not subscribed and used >= FREE_UNLOCKS:
        return {"success": False, "needs_payment": True, "amount": UNLOCK_PRICE, "free_unlocks_left": 0}
    await _do_unlock(user["id"], req.seeker_id, 0)
    left = 999 if subscribed else max(0, FREE_UNLOCKS - (used + 1))
    return {"success": True, "amount": 0, "free": True,
            "free_unlocks_left": left, "cv": redact_seeker(seeker, True)}


# ---------------- Payments (Razorpay + mock) ----------------
RAZORPAY_KEY_ID = os.environ.get("RAZORPAY_KEY_ID")
RAZORPAY_KEY_SECRET = os.environ.get("RAZORPAY_KEY_SECRET")


def razorpay_enabled() -> bool:
    return bool(RAZORPAY_KEY_ID and RAZORPAY_KEY_SECRET)


def _rzp_client():
    import razorpay
    return razorpay.Client(auth=(RAZORPAY_KEY_ID, RAZORPAY_KEY_SECRET))


class OrderIn(BaseModel):
    purpose: str  # "cv_unlock" | "subscription"
    seeker_id: Optional[str] = None


class VerifyIn(BaseModel):
    order_id: str
    purpose: str
    seeker_id: Optional[str] = None
    razorpay_payment_id: Optional[str] = None
    razorpay_signature: Optional[str] = None


def _amount_for(purpose: str) -> int:
    return SUBSCRIPTION_PRICE if purpose == "subscription" else UNLOCK_PRICE


@api_router.post("/pay/create-order")
async def create_order(body: OrderIn, user=Depends(get_current_user)):
    if body.purpose not in ("cv_unlock", "subscription"):
        raise HTTPException(status_code=400, detail="Invalid purpose")
    amount = _amount_for(body.purpose)
    now = datetime.now(timezone.utc).isoformat()
    if razorpay_enabled():
        order = await asyncio.to_thread(lambda: _rzp_client().order.create({
            "amount": amount * 100, "currency": "INR", "payment_capture": 1,
            "receipt": f"{body.purpose[:12]}-{uuid.uuid4().hex[:8]}"}))
        order_id, mock = order["id"], False
    else:
        order_id, mock = f"order_mock_{uuid.uuid4().hex[:14]}", True
    await db.payments.insert_one({"user_id": user["id"], "type": body.purpose, "amount": amount,
                                  "order_id": order_id, "seeker_id": body.seeker_id,
                                  "status": "created", "mock": mock, "at": now})
    provider = await db.providers.find_one({"user_id": user["id"]}, {"_id": 0}, sort=[("created_at", -1)]) or {}
    return {"order_id": order_id, "amount": amount, "amount_paise": amount * 100, "currency": "INR",
            "key_id": RAZORPAY_KEY_ID or "", "mock": mock, "purpose": body.purpose, "seeker_id": body.seeker_id,
            "prefill": {"name": provider.get("company_name", ""), "email": provider.get("email", ""),
                        "contact": user.get("mobile", "")}}


@api_router.post("/pay/verify")
async def verify_payment(body: VerifyIn, user=Depends(get_current_user)):
    rec = await db.payments.find_one({"order_id": body.order_id, "user_id": user["id"]}, {"_id": 0})
    if not rec:
        raise HTTPException(status_code=404, detail="Order not found")
    if razorpay_enabled() and not rec.get("mock"):
        try:
            await asyncio.to_thread(lambda: _rzp_client().utility.verify_payment_signature({
                "razorpay_order_id": body.order_id,
                "razorpay_payment_id": body.razorpay_payment_id,
                "razorpay_signature": body.razorpay_signature}))
        except Exception:
            raise HTTPException(status_code=400, detail="Payment verification failed")
    now = datetime.now(timezone.utc).isoformat()
    await db.payments.update_one({"order_id": body.order_id},
                                 {"$set": {"status": "paid", "paid_at": now,
                                           "payment_id": body.razorpay_payment_id}})
    if body.purpose == "cv_unlock":
        seeker = await db.seekers.find_one({"id": body.seeker_id}, {"_id": 0})
        if not seeker:
            raise HTTPException(status_code=404, detail="Profile not found")
        await _do_unlock(user["id"], body.seeker_id, UNLOCK_PRICE)
        return {"success": True, "purpose": "cv_unlock", "amount": UNLOCK_PRICE,
                "cv": redact_seeker(seeker, True)}
    expiry = (datetime.now(timezone.utc) + timedelta(days=365)).isoformat()
    await db.users.update_one({"id": user["id"]}, {"$set": {"subscribed": True, "subscription_expiry": expiry}})
    prov = await db.providers.find_one({"user_id": user["id"]}, {"_id": 0}, sort=[("created_at", -1)]) or {}
    base = round(SUBSCRIPTION_PRICE / 1.18, 2)
    inv = {"user_id": user["id"],
           "invoice_no": f"SC-{datetime.now(timezone.utc).strftime('%Y%m%d')}-{uuid.uuid4().hex[:6].upper()}",
           "amount": SUBSCRIPTION_PRICE, "base": base, "gst_pct": 18, "gst": round(SUBSCRIPTION_PRICE - base, 2),
           "company_name": prov.get("company_name", ""), "emailed_to": prov.get("email") or user.get("mobile", ""),
           "email_status": "simulated", "at": now}
    await db.invoices.insert_one(dict(inv))
    inv.pop("_id", None)
    logger.info(f"GST invoice {inv['invoice_no']} emailed (DEMO) to {inv['emailed_to']}")
    return {"success": True, "purpose": "subscription", "subscription_expiry": expiry,
            "amount": SUBSCRIPTION_PRICE, "invoice": inv}


# ---------------- Company subscription / banner / CV download ----------------
class BannerIn(BaseModel):
    title: str
    tagline: Optional[str] = ""
    image_url: Optional[str] = ""


@api_router.get("/company/status")
async def company_status(user=Depends(get_current_user)):
    sub = await is_subscribed(user["id"])
    used = await db.unlocks.count_documents({"user_id": user["id"]})
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    dl = await db.cv_downloads.count_documents({"user_id": user["id"], "day": today})
    banner = await db.banners.find_one({"user_id": user["id"]}, {"_id": 0})
    u = await db.users.find_one({"id": user["id"]}, {"_id": 0}) or {}
    return {"subscribed": sub, "subscription_expiry": u.get("subscription_expiry"),
            "free_unlocks_left": max(0, FREE_UNLOCKS - used), "unlock_price": UNLOCK_PRICE,
            "subscription_price": SUBSCRIPTION_PRICE, "downloads_today": dl, "download_limit": 20,
            "banner": banner}


@api_router.post("/company/banner")
async def set_banner(body: BannerIn, user=Depends(get_current_user)):
    if not await is_subscribed(user["id"]):
        raise HTTPException(status_code=403, detail="Home-page banner is part of the ₹1,99,000/yr plan")
    provider = await db.providers.find_one({"user_id": user["id"]}, {"_id": 0}, sort=[("created_at", -1)]) or {}
    await db.banners.update_one({"user_id": user["id"]},
        {"$set": {"user_id": user["id"], "title": body.title, "tagline": body.tagline,
                  "image_url": body.image_url, "company_name": provider.get("company_name", ""),
                  "updated_at": datetime.now(timezone.utc).isoformat()}}, upsert=True)
    return {"ok": True}


@api_router.get("/banners")
async def get_banners():
    rows = await db.banners.find({}, {"_id": 0}).to_list(50)
    out = [b for b in rows if await is_subscribed(b["user_id"])]
    return {"banners": out}


@api_router.get("/invoices")
async def my_invoices(user=Depends(get_current_user)):
    rows = await db.invoices.find({"user_id": user["id"]}, {"_id": 0}).sort("at", -1).to_list(100)
    return {"invoices": rows}


@api_router.post("/crew-needs")
async def create_crew_need(body: CrewNeedIn, user=Depends(get_current_user)):
    doc = body.model_dump()
    doc["id"] = str(uuid.uuid4())
    doc["user_id"] = user["id"]
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    await db.crew_needs.insert_one(dict(doc))
    doc.pop("_id", None)
    return doc


@api_router.get("/crew-needs")
async def list_crew_needs(user=Depends(get_current_user)):
    rows = await db.crew_needs.find({"user_id": user["id"]}, {"_id": 0}).sort("created_at", -1).to_list(500)
    return {"crew_needs": rows}


@api_router.delete("/crew-needs/{need_id}")
async def delete_crew_need(need_id: str, user=Depends(get_current_user)):
    res = await db.crew_needs.delete_one({"id": need_id, "user_id": user["id"]})
    if not res.deleted_count:
        raise HTTPException(status_code=404, detail="Crew requirement not found")
    return {"ok": True}


def build_seeker_cv_pdf(s: dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    import textwrap
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    w, h = A4
    c.setFillColorRGB(1, 0.35, 0)
    c.rect(0, h - 16 * mm, w, 16 * mm, fill=1, stroke=0)
    c.setFillColorRGB(1, 1, 1)
    c.setFont("Helvetica-Bold", 14)
    c.drawString(18 * mm, h - 11 * mm, "SMART CREWING — Standard CV")
    c.setFillColorRGB(0, 0, 0)
    y = h - 25 * mm
    c.setFont("Helvetica-Bold", 16)
    c.drawString(18 * mm, y, s.get("name", "Seafarer"))
    y -= 10 * mm
    c.setFont("Helvetica", 11)
    rows = [("Rank", s.get("rank")), ("Vessel Type", s.get("vessel_type")),
            ("Sea Time", s.get("sea_time")), ("COC", s.get("coc")),
            ("US Visa", s.get("us_visa")), ("Other Visa", s.get("other_visa", "")),
            ("Available From", s.get("availability")), ("Mobile", s.get("mobile")),
            ("Email", s.get("email"))]
    for k, v in rows:
        c.drawString(18 * mm, y, f"{k}: {v if v not in (None, '') else '—'}")
        y -= 7 * mm
    summ = s.get("summary", "")
    if summ:
        y -= 4 * mm
        c.setFont("Helvetica-Bold", 11); c.drawString(18 * mm, y, "Summary"); y -= 7 * mm
        c.setFont("Helvetica", 10)
        for line in textwrap.wrap(summ, 95):
            c.drawString(18 * mm, y, line); y -= 6 * mm
    c.setFont("Helvetica-Oblique", 8)
    c.drawString(18 * mm, 12 * mm, f"Generated by Smart Crewing · {datetime.now(timezone.utc).strftime('%Y-%m-%d')}")
    c.showPage(); c.save()
    return buf.getvalue()


@api_router.get("/cv/{seeker_id}/download")
async def download_cv(seeker_id: str, user=Depends(get_current_user)):
    if not await is_subscribed(user["id"]):
        raise HTTPException(status_code=403, detail="CV download (Standard format) is part of the ₹1,99,000/yr plan")
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    count = await db.cv_downloads.count_documents({"user_id": user["id"], "day": today})
    if count >= 20:
        raise HTTPException(status_code=429, detail="Daily limit of 20 CV downloads reached. Please try again tomorrow.")
    seeker = await db.seekers.find_one({"id": seeker_id}, {"_id": 0})
    if not seeker:
        raise HTTPException(status_code=404, detail="Profile not found")
    pdf = await asyncio.to_thread(build_seeker_cv_pdf, seeker)
    await db.cv_downloads.insert_one({"user_id": user["id"], "seeker_id": seeker_id, "day": today,
                                      "at": datetime.now(timezone.utc).isoformat()})
    fname = f"cv-{(seeker.get('name') or 'crew').replace(' ', '_')}.pdf"
    return StreamingResponse(io.BytesIO(pdf), media_type="application/pdf",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@api_router.post("/boost-profile")
async def boost_profile(user=Depends(get_current_user)):
    seeker = await db.seekers.find_one({"user_id": user["id"]}, {"_id": 0}, sort=[("created_at", -1)])
    if not seeker:
        raise HTTPException(status_code=404, detail="No profile to boost")
    await db.seekers.update_one({"id": seeker["id"]}, {"$set": {"boosted": True}})
    await db.payments.insert_one({"user_id": user["id"], "type": "boost", "amount": 100,
                                  "at": datetime.now(timezone.utc).isoformat()})
    fresh = await db.seekers.find_one({"id": seeker["id"]}, {"_id": 0})
    await notify_companies_of_seeker(fresh)
    return {"success": True, "amount": 100, "boosted": True}


# ---------------- Seeker eligibility (view company contacts) ----------------
class ShareIn(BaseModel):
    type: str  # group | contact


@api_router.get("/eligibility")
async def get_eligibility(user=Depends(get_current_user)):
    u = await db.users.find_one({"id": user["id"]}, {"_id": 0}) or {}
    return {
        "eligible": bool(u.get("reg_paid") or u.get("ref_completed")),
        "reg_paid": bool(u.get("reg_paid")),
        "ref_completed": bool(u.get("ref_completed")),
        "ref_signups": u.get("ref_signups", 0),
        "signup_target": SIGNUP_TARGET,
    }


@api_router.post("/referral/share")
async def referral_share(body: ShareIn, user=Depends(get_current_user)):
    if body.type not in ("group", "contact"):
        raise HTTPException(status_code=400, detail="type must be 'group' or 'contact'")
    field = "ref_group_shares" if body.type == "group" else "ref_contact_shares"
    await db.users.update_one({"id": user["id"]}, {"$inc": {field: 1}})
    u = await db.users.find_one({"id": user["id"]}, {"_id": 0})
    return {"group_shares": u.get("ref_group_shares", 0), "contact_shares": u.get("ref_contact_shares", 0)}


@api_router.post("/referral/claim")
async def referral_claim(user=Depends(get_current_user)):
    u = await db.users.find_one({"id": user["id"]}, {"_id": 0})
    g, c = u.get("ref_group_shares", 0), u.get("ref_contact_shares", 0)
    if g >= 2 or c >= 20:
        await db.users.update_one({"id": user["id"]}, {"$set": {"ref_completed": True}})
        await db.seekers.update_many({"user_id": user["id"]}, {"$set": {"boosted": True}})
        return {"success": True, "ref_completed": True}
    raise HTTPException(status_code=400, detail=f"Share with 2 groups OR 20 contacts first. Progress: {g}/2 groups, {c}/20 contacts.")


@api_router.post("/register/pay")
async def register_pay(user=Depends(get_current_user)):
    u = await db.users.find_one({"id": user["id"]}, {"_id": 0}) or {}
    if u.get("reg_paid"):
        return {"success": True, "amount": 100, "reg_paid": True, "already": True}
    await db.users.update_one({"id": user["id"]}, {"$set": {"reg_paid": True}})
    await db.payments.insert_one({"user_id": user["id"], "type": "registration", "amount": 100,
                                  "at": datetime.now(timezone.utc).isoformat()})
    return {"success": True, "amount": 100, "reg_paid": True}


@api_router.post("/get-noticed/pay")
async def get_noticed_pay(user=Depends(get_current_user)):
    """Mock ₹499 'Get noticed faster' unlock — promotes the seafarer's profile."""
    await db.users.update_one({"id": user["id"]}, {"$set": {"reg_paid": True}})
    await db.seekers.update_many({"user_id": user["id"]}, {"$set": {"boosted": True}})
    await db.payments.insert_one({"user_id": user["id"], "type": "get_noticed_faster", "amount": 499,
                                  "at": datetime.now(timezone.utc).isoformat()})
    for s in await db.seekers.find({"user_id": user["id"]}, {"_id": 0}).to_list(100):
        await notify_companies_of_seeker(s)
    return {"success": True, "amount": 499, "eligible": True}


@api_router.get("/promoted-companies")
async def promoted_companies(user=Depends(get_current_user)):
    u = await db.users.find_one({"id": user["id"]}, {"_id": 0}) or {}
    eligible = bool(u.get("reg_paid") or u.get("ref_completed"))
    seeker = await db.seekers.find_one({"user_id": user["id"]}, {"_id": 0})
    if not eligible or not seeker:
        return {"eligible": eligible, "companies": [], "count": 0}
    providers = await db.providers.find(ACTIVE_FILTER, {"_id": 0}).to_list(1000)
    matched = [p for p in providers if matches_requirement(seeker, p)]
    companies = [{"company_name": p.get("company_name"), "rank": p.get("rank"),
                  "vessel_type": p.get("vessel_type"), "shared_at": p.get("created_at", "")} for p in matched]
    return {"eligible": True, "companies": companies, "count": len(companies)}


@api_router.get("/dashboard")
async def dashboard(user=Depends(get_current_user)):
    seekers = await db.seekers.find({"user_id": user["id"]}, {"_id": 0}).to_list(100)
    providers = await db.providers.find({"user_id": user["id"]}, {"_id": 0}).to_list(100)
    unlocks = await db.unlocks.find({"user_id": user["id"]}, {"_id": 0}).to_list(100)
    payments = await db.payments.find({"user_id": user["id"]}, {"_id": 0}).to_list(100)
    docs_count = await db.documents.count_documents({"user_id": user["id"], "is_deleted": False})
    wa = await db.wa_alerts.find({"company_user_id": user["id"]}, {"_id": 0}).sort("created_at", -1).to_list(200)
    crew_needs = await db.crew_needs.find({"user_id": user["id"]}, {"_id": 0}).sort("created_at", -1).to_list(500)
    return {
        "seeker_profiles": seekers,
        "provider_requirements": providers,
        "unlocks": unlocks,
        "payments": payments,
        "documents_count": docs_count,
        "wa_alerts": wa,
        "crew_needs": crew_needs,
    }


@api_router.get("/wa-alerts")
async def get_wa_alerts(user=Depends(get_current_user)):
    rows = await db.wa_alerts.find({"company_user_id": user["id"]}, {"_id": 0}).sort("created_at", -1).to_list(200)
    return {"alerts": rows, "demo": not twilio_enabled()}


# ---------------- DPDP data rights ----------------
class DobIn(BaseModel):
    dob: str  # DD/MM/YYYY


def _compute_age(dob: str):
    m = re.match(r"^\s*(\d{1,2})[/-](\d{1,2})[/-](\d{4})\s*$", dob or "")
    if not m:
        return None
    d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    try:
        birth = datetime(y, mo, d, tzinfo=timezone.utc)
    except ValueError:
        return None
    today = datetime.now(timezone.utc)
    return today.year - birth.year - ((today.month, today.day) < (birth.month, birth.day))


@api_router.post("/me/dob")
async def set_dob(body: DobIn, user=Depends(get_current_user)):
    age = _compute_age(body.dob)
    if age is None or age < 0 or age > 120:
        raise HTTPException(status_code=400, detail="Please enter a valid date (DD/MM/YYYY).")
    if age < 18:
        raise HTTPException(status_code=403, detail="You must be 18 or older to use Smart Crewing.")
    await db.users.update_one({"id": user["id"]}, {"$set": {"dob": body.dob, "is_adult": True}})
    return {"ok": True, "age": age}


@api_router.get("/me/export")
async def export_my_data(user=Depends(get_current_user)):
    """DPDP right to access — a machine-readable copy of the user's personal data."""
    u = await db.users.find_one({"id": user["id"]}, {"_id": 0}) or {}
    seekers = await db.seekers.find({"user_id": user["id"]}, {"_id": 0}).to_list(100)
    providers = await db.providers.find({"user_id": user["id"]}, {"_id": 0}).to_list(100)
    documents = await db.documents.find({"user_id": user["id"], "is_deleted": False}, {"_id": 0, "data": 0}).to_list(200)
    unlocks = await db.unlocks.find({"user_id": user["id"]}, {"_id": 0}).to_list(500)
    payments = await db.payments.find({"user_id": user["id"]}, {"_id": 0}).to_list(500)
    return {
        "exported_at": datetime.now(timezone.utc).isoformat(),
        "account": {"mobile": u.get("mobile"), "role": u.get("role"), "dob": u.get("dob"),
                    "consent_given": u.get("consent_given"), "consent_at": u.get("consent_at"),
                    "consent_version": u.get("consent_version")},
        "profiles": seekers, "requirements": providers, "documents": documents,
        "unlocks": unlocks, "payments": payments,
    }


@api_router.post("/me/request-deletion")
async def request_deletion(user=Depends(get_current_user)):
    """DPDP right to erasure — records are withdrawn from all public listings and moved to an
    admin-only 'Requested for deletion under DPDP' folder (not hard-deleted)."""
    now = datetime.now(timezone.utc).isoformat()
    flag = {"deletion_requested": True, "deletion_requested_at": now}
    await db.users.update_one({"id": user["id"]}, {"$set": flag})
    await db.seekers.update_many({"user_id": user["id"]}, {"$set": flag})
    await db.providers.update_many({"user_id": user["id"]}, {"$set": flag})
    await db.documents.update_many({"user_id": user["id"]}, {"$set": flag})
    await db.deletion_requests.update_one(
        {"user_id": user["id"]},
        {"$set": {"user_id": user["id"], "mobile": user.get("mobile"), "requested_at": now}},
        upsert=True)
    return {"ok": True, "folder": "Requested for deletion under DPDP"}


@api_router.get("/grievance")
async def grievance():
    return {"grievance_officer": GRIEVANCE_OFFICER, "consent_version": DPDP_CONSENT_VERSION}


@api_router.get("/admin/deletion-requests")
async def admin_deletion_requests(admin=Depends(require_admin)):
    users = await db.users.find({"deletion_requested": True}, {"_id": 0}).sort("deletion_requested_at", -1).to_list(500)
    out = []
    for u in users:
        seekers = await db.seekers.find({"user_id": u["id"]}, {"_id": 0}).to_list(50)
        providers = await db.providers.find({"user_id": u["id"]}, {"_id": 0}).to_list(50)
        docs = await db.documents.count_documents({"user_id": u["id"]})
        out.append({"mobile": u.get("mobile"), "requested_at": u.get("deletion_requested_at"),
                    "profiles": seekers, "requirements": providers, "documents_count": docs})
    return {"folder": "Requested for deletion under DPDP", "count": len(out), "requests": out}


@api_router.post("/feedback")
async def submit_feedback(fb: FeedbackIn, user=Depends(get_current_user)):
    doc = fb.model_dump()
    doc["id"] = str(uuid.uuid4())
    doc["user_id"] = user["id"]
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    await db.feedback.insert_one(dict(doc))
    return {"success": True}


# ---------------- Documents ----------------
CREW_LIMIT = 100 * 1024 * 1024   # 100 MB for seafarer/crew users
COMPANY_LIMIT = 5 * 1024 * 1024  # 5 MB for hiring companies


async def user_storage_used(user_id: str) -> int:
    pipeline = [
        {"$match": {"user_id": user_id, "is_deleted": False}},
        {"$group": {"_id": None, "total": {"$sum": "$size"}}},
    ]
    agg = await db.documents.aggregate(pipeline).to_list(1)
    return agg[0]["total"] if agg else 0


async def user_storage_limit(user_id: str) -> int:
    has_provider = await db.providers.count_documents({"user_id": user_id}) > 0
    has_seeker = await db.seekers.count_documents({"user_id": user_id}) > 0
    if has_provider and not has_seeker:
        return COMPANY_LIMIT
    return CREW_LIMIT


@api_router.get("/documents/usage")
async def documents_usage(user=Depends(get_current_user)):
    used = await user_storage_used(user["id"])
    limit = await user_storage_limit(user["id"])
    utype = "company" if limit == COMPANY_LIMIT else "crew"
    return {"used": used, "limit": limit, "remaining": max(0, limit - used), "user_type": utype}


@api_router.post("/documents")
async def upload_document(file: UploadFile = File(...), category: str = Query("Other"),
                          user=Depends(get_current_user)):
    ext = file.filename.split(".")[-1].lower() if "." in file.filename else "bin"
    data = await file.read()
    size = len(data)
    if size > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Each file must be 5 MB or smaller. Please compress your document (PDF recommended).")
    used = await user_storage_used(user["id"])
    limit = await user_storage_limit(user["id"])
    if used + size > limit:
        remaining_mb = max(0, limit - used) / (1024 * 1024)
        limit_mb = round(limit / (1024 * 1024))
        raise HTTPException(
            status_code=400,
            detail=f"Storage limit of {limit_mb} MB reached. You have {remaining_mb:.1f} MB free.",
        )
    path = f"{APP_NAME}/uploads/{user['id']}/{uuid.uuid4()}.{ext}"
    content_type = file.content_type or MIME_TYPES.get(ext, "application/octet-stream")
    result = put_object(path, data, content_type)
    doc = {
        "id": str(uuid.uuid4()),
        "user_id": user["id"],
        "storage_path": result["path"],
        "original_filename": file.filename,
        "category": category,
        "content_type": content_type,
        "size": result.get("size", size),
        "is_deleted": False,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.documents.insert_one(dict(doc))
    doc.pop("_id", None)
    return doc


@api_router.get("/documents")
async def list_documents(user=Depends(get_current_user)):
    docs = await db.documents.find({"user_id": user["id"], "is_deleted": False}, {"_id": 0}).to_list(1000)
    return docs


@api_router.get("/documents/{doc_id}/download")
async def download_document(doc_id: str, authorization: Optional[str] = Header(None),
                            auth: Optional[str] = Query(None)):
    token = None
    if authorization and authorization.startswith("Bearer "):
        token = authorization.split(" ", 1)[1]
    elif auth:
        token = auth
    if not token:
        raise HTTPException(status_code=401, detail="Not authenticated")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=["HS256"])
    except jwt.PyJWTError:
        raise HTTPException(status_code=401, detail="Invalid token")
    record = await db.documents.find_one({"id": doc_id, "is_deleted": False}, {"_id": 0})
    if not record or record["user_id"] != payload["user_id"]:
        raise HTTPException(status_code=404, detail="Document not found")
    data, content_type = get_object(record["storage_path"])
    return Response(content=data, media_type=record.get("content_type", content_type),
                    headers={"Content-Disposition": f'inline; filename="{record["original_filename"]}"'})


@api_router.delete("/documents/{doc_id}")
async def delete_document(doc_id: str, user=Depends(get_current_user)):
    record = await db.documents.find_one({"id": doc_id}, {"_id": 0})
    if not record or record["user_id"] != user["id"]:
        raise HTTPException(status_code=404, detail="Document not found")
    await db.documents.update_one({"id": doc_id}, {"$set": {"is_deleted": True}})
    return {"success": True}


# ---------------- CV Parsing → Standard Maritime CV ----------------
STANDARD_CV_TEMPLATE = {
    "personal": {"full_name": "", "date_of_birth": "", "nationality": "", "gender": "", "marital_status": "", "place_of_birth": ""},
    "contact": {"mobile": "", "email": "", "address": "", "nearest_airport": ""},
    "professional": {
        "present_rank": "", "applied_rank": "", "total_sea_experience_months": "",
        "coc_grade": "", "coc_issuing_authority": "", "coc_number": "",
        "indos_number": "", "passport_number": "", "passport_expiry": "",
        "cdc_number": "", "cdc_expiry": "",
    },
    "visas": {"us_visa_status": "", "schengen_visa": "", "other_visas": ""},
    "availability": {"available_from": "", "expected_wages": ""},
    "sea_service": [],  # {company, vessel_name, vessel_type, rank, grt_bhp, from_date, to_date, months}
    "certificates": [],  # {name, number, issue_date, expiry_date}
    "next_of_kin": {"name": "", "relation": "", "contact": ""},
}

REQUIRED_PATHS = [
    "personal.full_name", "personal.nationality", "contact.mobile", "contact.email",
    "professional.present_rank", "professional.total_sea_experience_months",
    "professional.coc_grade", "availability.available_from",
]


def extract_text_from_bytes(data: bytes, filename: str) -> Optional[str]:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        if ext == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((p.extract_text() or "") for p in reader.pages).strip()
        if ext == "docx":
            from docx import Document as DocxDocument
            d = DocxDocument(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts).strip()
        if ext in ("txt", "csv"):
            return data.decode("utf-8", errors="ignore").strip()
    except Exception as e:
        logger.warning(f"text extraction failed: {e}")
    return None


CV_SYSTEM_PROMPT = (
    "You are a maritime crewing data-extraction engine. Extract information from a seafarer's "
    "biodata / CV and return it STRICTLY as JSON matching this exact schema (keys and nesting):\n"
    + json.dumps(STANDARD_CV_TEMPLATE)
    + "\nRules: (1) Fill every field you can find. Use \"\" for anything missing or illegible. "
    "(2) total_sea_experience_months must be a number as string. (3) Dates strictly as DD/MM/YYYY (day/month/year with slashes, e.g. 15/08/2026). "
    "(4) sea_service and certificates are arrays of objects with the shown keys. "
    "(5) Add an extra top-level key \"uncertain_fields\": a list of dotted paths (e.g. \"personal.full_name\", "
    "\"sea_service[0].vessel_name\") whose values you GUESSED or found hard to read (common with handwriting). "
    "(6) Respond with ONLY the JSON object, no prose, no code fences."
)


def _clean_json(text: str) -> dict:
    text = text.strip()
    if text.startswith("```"):
        text = text.split("```", 2)[1] if "```" in text[3:] else text
        text = text.replace("json", "", 1).strip("` \n")
    start, end = text.find("{"), text.rfind("}")
    if start != -1 and end != -1:
        text = text[start:end + 1]
    return json.loads(text)


async def parse_cv_with_ai(text: Optional[str], image_b64: Optional[str]) -> dict:
    from emergentintegrations.llm.chat import LlmChat, UserMessage, ImageContent
    chat = LlmChat(api_key=EMERGENT_LLM_KEY, session_id=str(uuid.uuid4()),
                   system_message=CV_SYSTEM_PROMPT).with_model("openai", "gpt-5.4")
    if image_b64:
        msg = UserMessage(text="Extract this seafarer CV into the JSON schema. It may be scanned or handwritten.",
                          file_contents=[ImageContent(image_base64=image_b64)])
    else:
        msg = UserMessage(text="Extract this CV text into the JSON schema:\n\n" + (text or ""))
    resp = await chat.send_message(msg)
    return _clean_json(resp if isinstance(resp, str) else str(resp))


SEA_KEYS = ["company", "vessel_name", "vessel_type", "rank", "grt_bhp", "from_date", "to_date", "months"]
CERT_KEYS = ["name", "number", "issue_date", "expiry_date"]
SEA_ALIASES = {"from_date": ["from", "sign_on", "signon", "join_date"], "to_date": ["to", "sign_off", "signoff", "release_date"], "months": ["duration"]}


def _coerce_rows(rows, keys, aliases):
    out = []
    for row in rows or []:
        if not isinstance(row, dict):
            continue
        r = {}
        for k in keys:
            v = row.get(k)
            if (v is None or v == "") and k in aliases:
                for a in aliases[k]:
                    if row.get(a):
                        v = row.get(a)
                        break
            r[k] = str(v or "")
        out.append(r)
    return out


def merge_template(parsed: dict) -> dict:
    out = json.loads(json.dumps(STANDARD_CV_TEMPLATE))
    for section, val in out.items():
        if isinstance(val, dict):
            for k in val:
                out[section][k] = str(parsed.get(section, {}).get(k, "") or "")
    out["sea_service"] = _coerce_rows(parsed.get("sea_service"), SEA_KEYS, SEA_ALIASES)
    out["certificates"] = _coerce_rows(parsed.get("certificates"), CERT_KEYS, {})
    return out


def _path_exists(data: dict, path: str) -> bool:
    cur = data
    for name, idx in re.findall(r'([a-zA-Z_]+)(\[\d+\])?', path):
        if not isinstance(cur, dict) or name not in cur:
            return False
        cur = cur[name]
        if idx:
            i = int(idx[1:-1])
            if not isinstance(cur, list) or i >= len(cur):
                return False
            cur = cur[i]
    return True


def normalize_uncertain(paths, data: dict):
    out = []
    for p in paths or []:
        if not isinstance(p, str):
            continue
        p = re.sub(r'\.from(?![_a-zA-Z])', '.from_date', p)
        p = re.sub(r'\.to(?![_a-zA-Z])', '.to_date', p)
        if _path_exists(data, p) and p not in out:
            out.append(p)
    return out


_FULL_DATE_FMTS = ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%Y/%m/%d",
                   "%d %m %Y", "%d %b %Y", "%d %B %Y", "%d-%b-%Y", "%d-%b-%y", "%d %b %y")
_MY_DATE_FMTS = ("%m/%Y", "%m-%Y", "%b %Y", "%B %Y", "%b-%y", "%B-%Y", "%m %Y")


def normalize_date(s):
    if not s or not isinstance(s, str):
        return s
    t = s.strip()
    if re.fullmatch(r'\d{4}', t):  # year only
        return t
    for fmt in _FULL_DATE_FMTS:
        try:
            dt = datetime.strptime(t, fmt)
            return f"{dt.day:02d}/{dt.month:02d}/{dt.year}"
        except ValueError:
            pass
    for fmt in _MY_DATE_FMTS:
        try:
            dt = datetime.strptime(t, fmt)
            return f"{dt.month:02d}/{dt.year}"
        except ValueError:
            pass
    return s  # leave keywords like "Immediate" untouched


def apply_dates(data: dict) -> dict:
    for section, key in [("personal", "date_of_birth"), ("professional", "passport_expiry"),
                         ("professional", "cdc_expiry"), ("availability", "available_from")]:
        if isinstance(data.get(section), dict) and key in data[section]:
            data[section][key] = normalize_date(data[section][key])
    for row in data.get("sea_service", []) or []:
        for k in ("from_date", "to_date"):
            row[k] = normalize_date(row.get(k, ""))
    for row in data.get("certificates", []) or []:
        for k in ("issue_date", "expiry_date"):
            row[k] = normalize_date(row.get(k, ""))
    return data


def get_path(data: dict, path: str):
    cur = data
    for part in path.split("."):
        cur = cur.get(part, "") if isinstance(cur, dict) else ""
    return cur


def compute_missing(data: dict) -> List[str]:
    return [p for p in REQUIRED_PATHS if not str(get_path(data, p)).strip()]


class ParseRequest(BaseModel):
    document_id: str


class CVUpdate(BaseModel):
    data: dict
    uncertain_fields: Optional[List[str]] = None


async def _store_cv(user_id: str, document_id: str, filename: str, data: dict, uncertain: List[str]) -> dict:
    doc = {
        "id": str(uuid.uuid4()),
        "user_id": user_id,
        "document_id": document_id,
        "source_filename": filename,
        "data": data,
        "uncertain_fields": uncertain,
        "missing_fields": compute_missing(data),
        "status": "parsed",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.parsed_cvs.insert_one(dict(doc))
    doc.pop("_id", None)
    return doc


@api_router.post("/cv/parse")
async def parse_cv(req: ParseRequest, user=Depends(get_current_user)):
    record = await db.documents.find_one({"id": req.document_id, "is_deleted": False}, {"_id": 0})
    if not record or record["user_id"] != user["id"]:
        raise HTTPException(status_code=404, detail="Document not found")
    raw, _ = get_object(record["storage_path"])
    ext = record["original_filename"].rsplit(".", 1)[-1].lower()
    image_b64, text = None, None
    if ext in ("png", "jpg", "jpeg", "webp", "gif"):
        image_b64 = base64.b64encode(raw).decode()
    else:
        text = extract_text_from_bytes(raw, record["original_filename"])
        if not text:
            raise HTTPException(status_code=400, detail="Could not read text from this file. For scanned or handwritten CVs, upload an image (PNG/JPG).")
    try:
        parsed = await parse_cv_with_ai(text, image_b64)
    except Exception as e:
        logger.error(f"CV parse failed: {e}")
        raise HTTPException(status_code=502, detail="AI parsing failed. Please try again.")
    uncertain = parsed.pop("uncertain_fields", []) if isinstance(parsed, dict) else []
    data = merge_template(parsed)
    apply_dates(data)
    uncertain = normalize_uncertain(uncertain, data)
    result = await _store_cv(user["id"], req.document_id, record["original_filename"], data, uncertain)
    return result


@api_router.get("/cv")
async def get_my_cvs(user=Depends(get_current_user)):
    items = await db.parsed_cvs.find({"user_id": user["id"]}, {"_id": 0}).sort("created_at", -1).to_list(50)
    return items


@api_router.put("/cv/{cv_id}")
async def update_cv(cv_id: str, upd: CVUpdate, user=Depends(get_current_user)):
    cv = await db.parsed_cvs.find_one({"id": cv_id}, {"_id": 0})
    if not cv:
        raise HTTPException(status_code=404, detail="CV not found")
    if cv["user_id"] != user["id"] and not user.get("is_admin"):
        raise HTTPException(status_code=403, detail="Not allowed")
    data = merge_template(upd.data)
    apply_dates(data)
    update = {
        "data": data,
        "missing_fields": compute_missing(data),
        "status": "edited",
        "updated_at": datetime.now(timezone.utc).isoformat(),
        "edited_by": "admin" if user.get("is_admin") and cv["user_id"] != user["id"] else "user",
    }
    if upd.uncertain_fields is not None:
        update["uncertain_fields"] = upd.uncertain_fields
    await db.parsed_cvs.update_one({"id": cv_id}, {"$set": update})
    cv.update(update)
    return cv


@api_router.get("/cv/{cv_id}/pdf")
async def cv_pdf(cv_id: str, authorization: Optional[str] = Header(None), auth: Optional[str] = Query(None)):
    token = None
    if authorization and authorization.startswith("Bearer "):
        token = authorization.split(" ", 1)[1]
    elif auth:
        token = auth
    if not token:
        raise HTTPException(status_code=401, detail="Not authenticated")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=["HS256"])
    except jwt.PyJWTError:
        raise HTTPException(status_code=401, detail="Invalid token")
    cv = await db.parsed_cvs.find_one({"id": cv_id}, {"_id": 0})
    if not cv:
        raise HTTPException(status_code=404, detail="CV not found")
    is_admin = payload.get("mobile") in ADMIN_MOBILES
    if cv["user_id"] != payload["user_id"] and not is_admin:
        raise HTTPException(status_code=403, detail="Not allowed")
    pdf_bytes = build_cv_pdf(cv["data"])
    fname = (cv["data"].get("personal", {}).get("full_name") or "seafarer").replace(" ", "_")
    return Response(content=pdf_bytes, media_type="application/pdf",
                    headers={"Content-Disposition": f'inline; filename="{fname}_CV.pdf"'})


def build_cv_pdf(data: dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm,
                            leftMargin=16 * mm, rightMargin=16 * mm)
    styles = getSampleStyleSheet()
    orange = colors.HexColor("#FF5A00")
    h = ParagraphStyle("h", parent=styles["Heading2"], textColor=orange, fontSize=12, spaceBefore=10, spaceAfter=4)
    title = ParagraphStyle("t", parent=styles["Title"], fontSize=20, textColor=colors.HexColor("#0A0B12"))
    small = ParagraphStyle("s", parent=styles["Normal"], fontSize=9)
    elems = []

    p = data.get("personal", {})
    elems.append(Paragraph(p.get("full_name") or "Seafarer CV", title))
    elems.append(Paragraph("STANDARD MARITIME CV — Smart Crewing", small))
    elems.append(Spacer(1, 8))

    def kv_table(pairs):
        rows = [[Paragraph(f"<b>{k}</b>", small), Paragraph(str(v or "—"), small)] for k, v in pairs]
        t = Table(rows, colWidths=[55 * mm, 110 * mm])
        t.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#DDDDDD")),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return t

    elems.append(Paragraph("Personal Details", h))
    elems.append(kv_table([("Full Name", p.get("full_name")), ("Date of Birth", p.get("date_of_birth")),
                           ("Nationality", p.get("nationality")), ("Gender", p.get("gender")),
                           ("Marital Status", p.get("marital_status")), ("Place of Birth", p.get("place_of_birth"))]))
    c = data.get("contact", {})
    elems.append(Paragraph("Contact", h))
    elems.append(kv_table([("Mobile", c.get("mobile")), ("Email", c.get("email")),
                           ("Address", c.get("address")), ("Nearest Airport", c.get("nearest_airport"))]))
    pr = data.get("professional", {})
    elems.append(Paragraph("Professional & Certification", h))
    elems.append(kv_table([("Present Rank", pr.get("present_rank")), ("Applied Rank", pr.get("applied_rank")),
                           ("Total Sea Experience (months)", pr.get("total_sea_experience_months")),
                           ("COC Grade", pr.get("coc_grade")), ("COC Authority", pr.get("coc_issuing_authority")),
                           ("COC Number", pr.get("coc_number")), ("INDOS", pr.get("indos_number")),
                           ("Passport", pr.get("passport_number")), ("Passport Expiry", pr.get("passport_expiry")),
                           ("CDC Number", pr.get("cdc_number")), ("CDC Expiry", pr.get("cdc_expiry"))]))
    v = data.get("visas", {})
    av = data.get("availability", {})
    elems.append(Paragraph("Visas & Availability", h))
    elems.append(kv_table([("US Visa", v.get("us_visa_status")), ("Schengen", v.get("schengen_visa")),
                           ("Other Visas", v.get("other_visas")), ("Available From", av.get("available_from")),
                           ("Expected Wages", av.get("expected_wages"))]))

    nok = data.get("next_of_kin", {})
    if any((nok or {}).values()):
        elems.append(Paragraph("Next of Kin", h))
        elems.append(kv_table([("Name", nok.get("name")), ("Relation", nok.get("relation")), ("Contact", nok.get("contact"))]))

    ss = data.get("sea_service", []) or []
    if ss:
        elems.append(Paragraph("Sea Service History", h))
        head = ["Company", "Vessel", "Type", "Rank", "GRT/BHP", "From", "To"]
        rows = [[Paragraph(f"<b>{x}</b>", small) for x in head]]
        for s in ss:
            rows.append([Paragraph(str(s.get(k, "") or "—"), small) for k in
                        ["company", "vessel_name", "vessel_type", "rank", "grt_bhp", "from_date", "to_date"]])
        t = Table(rows, colWidths=[28 * mm, 26 * mm, 22 * mm, 24 * mm, 20 * mm, 22 * mm, 22 * mm])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), orange), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CCCCCC")),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        elems.append(t)

    certs = data.get("certificates", []) or []
    if certs:
        elems.append(Paragraph("Certificates & Licenses", h))
        head = ["Certificate", "Number", "Issue", "Expiry"]
        rows = [[Paragraph(f"<b>{x}</b>", small) for x in head]]
        for cc in certs:
            rows.append([Paragraph(str(cc.get(k, "") or "—"), small) for k in ["name", "number", "issue_date", "expiry_date"]])
        t = Table(rows, colWidths=[70 * mm, 40 * mm, 27 * mm, 27 * mm])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), orange), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CCCCCC")),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        elems.append(t)

    doc.build(elems)
    return buf.getvalue()


# ---------------- Admin ----------------
@api_router.get("/admin/feedback")
async def admin_list_feedback(admin=Depends(require_admin)):
    items = await db.feedback.find({}, {"_id": 0}).sort("created_at", -1).to_list(500)
    for it in items:
        u = await db.users.find_one({"id": it.get("user_id")}, {"_id": 0})
        it["mobile"] = u.get("mobile") if u else ""
    return items


@api_router.get("/admin/cvs")
async def admin_list_cvs(admin=Depends(require_admin)):
    items = await db.parsed_cvs.find({}, {"_id": 0}).sort("updated_at", -1).to_list(500)
    return items


@api_router.get("/admin/stats")
async def admin_stats(admin=Depends(require_admin)):
    return {
        "users": await db.users.count_documents({}),
        "seekers": await db.seekers.count_documents({}),
        "providers": await db.providers.count_documents({}),
        "documents": await db.documents.count_documents({"is_deleted": False}),
        "parsed_cvs": await db.parsed_cvs.count_documents({}),
        "feedback": await db.feedback.count_documents({}),
        "deletion_requests": await db.users.count_documents({"deletion_requested": True}),
    }


# ---------------- Movement / Availability ----------------
class MovementIn(BaseModel):
    status: str  # Onboard | On Leave | Available Now | Available From
    available_from: Optional[str] = ""
    vessel: Optional[str] = ""
    note: Optional[str] = ""


@api_router.post("/movement")
async def add_movement(m: MovementIn, user=Depends(get_current_user)):
    doc = m.model_dump()
    doc["id"] = str(uuid.uuid4())
    doc["user_id"] = user["id"]
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    await db.movements.insert_one(dict(doc))
    await db.users.update_one({"id": user["id"]}, {"$set": {
        "current_status": m.status, "available_from": m.available_from, "current_vessel": m.vessel,
    }})
    doc.pop("_id", None)
    return doc


@api_router.get("/movement")
async def list_movements(user=Depends(get_current_user)):
    items = await db.movements.find({"user_id": user["id"]}, {"_id": 0}).sort("created_at", -1).to_list(200)
    latest = await db.users.find_one({"id": user["id"]}, {"_id": 0})
    return {"history": items, "current": {
        "status": latest.get("current_status", ""),
        "available_from": latest.get("available_from", ""),
        "vessel": latest.get("current_vessel", ""),
    }}


@api_router.get("/")
async def root():
    return {"message": "Smart Crewing API"}


# ---------------- Seed sample data ----------------
SAMPLE_SEEKERS = [
    {"rank": "Chief Engineer", "vessel_type": "Oil Tanker", "sea_time": "30+ months", "coc": "India",
     "us_visa": "Valid", "other_visa": "Schengen", "name": "Rajesh Nair", "mobile": "+919812345670",
     "email": "rajesh.nair@example.com", "availability": "Immediate", "boosted": True},
    {"rank": "Chief Engineer", "vessel_type": "Oil Tanker", "sea_time": "18–30 months", "coc": "Singapore",
     "us_visa": "Not Available", "other_visa": "", "name": "Wei Chen", "mobile": "+6591234567",
     "email": "wei.chen@example.com", "availability": "01/07/2026", "boosted": False},
    {"rank": "Captain / Master", "vessel_type": "Container", "sea_time": "30+ months", "coc": "UK",
     "us_visa": "Valid", "other_visa": "UK", "name": "James Whitfield", "mobile": "+447700900123",
     "email": "james.w@example.com", "availability": "Immediate", "boosted": False},
    {"rank": "Second Officer", "vessel_type": "Bulk Carrier", "sea_time": "9–18 months", "coc": "Indonesia",
     "us_visa": "Not Valid", "other_visa": "", "name": "Budi Santoso", "mobile": "+628123456789",
     "email": "budi.s@example.com", "availability": "15/07/2026", "boosted": True},
    {"rank": "Able Seaman (AB)", "vessel_type": "Container", "sea_time": "18–30 months", "coc": "Philippines",
     "us_visa": "Valid", "other_visa": "", "name": "Mario Reyes", "mobile": "+639171234567",
     "email": "mario.r@example.com", "availability": "Immediate", "boosted": False},
    {"rank": "Third Engineer", "vessel_type": "LNG / LPG Tanker", "sea_time": "9–18 months", "coc": "India",
     "us_visa": "Not Available", "other_visa": "", "name": "Arjun Menon", "mobile": "+919876500011",
     "email": "arjun.m@example.com", "availability": "10/08/2026", "boosted": False},
]

SAMPLE_PROVIDERS = [
    {"rank": "Chief Engineer", "vessel_type": "Oil Tanker", "sea_time": "18–30 months",
     "coc": ["India", "Singapore", "UK"], "us_visa": "Preferred", "company_name": "Oceanic Tankers Ltd",
     "email": "crewing@oceanictankers.com", "availability": "Immediate"},
    {"rank": "Captain / Master", "vessel_type": "Container", "sea_time": "30+ months",
     "coc": ["Any"], "us_visa": "Must Have", "company_name": "BlueLine Container Co",
     "email": "hr@blueline.com", "availability": "01/08/2026"},
    {"rank": "Second Officer", "vessel_type": "Bulk Carrier", "sea_time": "0–9 months",
     "coc": ["Indonesia", "India", "Any"], "us_visa": "Not Required", "company_name": "Global Bulk Shipping",
     "email": "careers@globalbulk.com", "availability": "Immediate"},
]


@app.on_event("startup")
async def seed_data():
    try:
        init_storage()
        logger.info("Storage initialized")
    except Exception as e:
        logger.error(f"Storage init failed: {e}")
    if await db.seekers.count_documents({"seed": True}) == 0:
        for s in SAMPLE_SEEKERS:
            doc = dict(s)
            doc.update({
                "id": str(uuid.uuid4()), "user_id": "seed", "role": "seeker", "seed": True,
                "consent": True, "summary": build_template_summary("seeker", s),
                "referral_code": uuid.uuid4().hex[:8],
                "created_at": datetime.now(timezone.utc).isoformat(),
            })
            await db.seekers.insert_one(doc)
    if await db.providers.count_documents({"seed": True}) == 0:
        for p in SAMPLE_PROVIDERS:
            doc = dict(p)
            doc.update({
                "id": str(uuid.uuid4()), "user_id": "seed", "role": "provider", "seed": True,
                "consent": True, "summary": build_template_summary("provider", p),
                "referral_code": uuid.uuid4().hex[:8],
                "created_at": datetime.now(timezone.utc).isoformat(),
            })
            await db.providers.insert_one(doc)
    logger.info("Seed data ensured")
    # backfill canonical mobile_key on legacy users so they log in with/without country code
    async for u in db.users.find({"mobile_key": {"$exists": False}}, {"id": 1, "mobile": 1}):
        await db.users.update_one({"id": u["id"]}, {"$set": {"mobile_key": mobile_key(u.get("mobile", ""))}})
    # migrate stored CVs to DD MM YYYY date format
    async for cv in db.parsed_cvs.find({}):
        data = cv.get("data", {})
        before = json.dumps(data, sort_keys=True)
        apply_dates(data)
        if json.dumps(data, sort_keys=True) != before:
            await db.parsed_cvs.update_one({"id": cv["id"]}, {"$set": {"data": data}})
    logger.info("CV date migration ensured")


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
